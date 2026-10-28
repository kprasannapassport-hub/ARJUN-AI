"""
Arjun's TOOLS - ordinary Python functions the AI is allowed to use.
The text inside each triple-quoted block tells the AI what the function does,
so it must be clear.
"""
import os
import re
import json
import time
import datetime
import webbrowser
from urllib.parse import quote

import config
import history


# ======================================================================
# TO-DO LIST
# ======================================================================
TASKS_FILE = config.DATA_DIR / "tasks.json"


def _load_json(path, default):
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return default


def _save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def add_task(title: str, due: str = "") -> str:
    """Add a new task to the to-do list.

    Args:
        title: What needs to be done.
        due: Optional due date or time as text, for example '2026-11-02' or 'tomorrow 5pm'.
    """
    tasks = _load_json(TASKS_FILE, [])
    new_id = max((t["id"] for t in tasks), default=0) + 1
    tasks.append({"id": new_id, "title": title, "due": due, "done": False})
    _save_json(TASKS_FILE, tasks)
    return f"Added task #{new_id}: {title}" + (f" (due {due})" if due else "")


def list_tasks(show_done: bool = False) -> str:
    """Show the to-do list.

    Args:
        show_done: Set true to also include tasks already completed.
    """
    tasks = [t for t in _load_json(TASKS_FILE, []) if show_done or not t["done"]]
    if not tasks:
        return "No tasks."
    lines = []
    for t in tasks:
        mark = "[x]" if t["done"] else "[ ]"
        due = f" (due {t['due']})" if t["due"] else ""
        lines.append(f"{mark} #{t['id']} {t['title']}{due}")
    return "\n".join(lines)


def complete_task(task_id: int) -> str:
    """Mark a task as done.

    Args:
        task_id: The number of the task, as shown in the list.
    """
    tasks = _load_json(TASKS_FILE, [])
    for t in tasks:
        if t["id"] == int(task_id):
            t["done"] = True
            _save_json(TASKS_FILE, tasks)
            return f"Marked #{task_id} as done."
    return f"No task with id {task_id}."


def delete_task(task_id: int) -> str:
    """Delete a task from the list.

    Args:
        task_id: The number of the task, as shown in the list.
    """
    tasks = _load_json(TASKS_FILE, [])
    remaining = [t for t in tasks if t["id"] != int(task_id)]
    if len(remaining) == len(tasks):
        return f"No task with id {task_id}."
    _save_json(TASKS_FILE, remaining)
    return f"Deleted task #{task_id}."


# ======================================================================
# WEB SEARCH + TIME
# ======================================================================
def search_web(query: str) -> str:
    """Search the internet for current information and return the top results.

    Args:
        query: What to search for.
    """
    try:
        from ddgs import DDGS
        results = DDGS().text(query, max_results=5)
    except Exception as e:
        return f"Search failed: {e}"
    if not results:
        return "No results found."
    return "\n\n".join(f"{r['title']}\n{r['body']}\n{r['href']}" for r in results)


def get_time() -> str:
    """Get the current date and time."""
    return datetime.datetime.now().strftime("%A, %d %B %Y, %I:%M %p")


# ======================================================================
# HISTORY SEARCH
# ======================================================================
def search_history(query: str, limit: int = 8) -> str:
    """Search through everything the user and Arjun said in past conversations.

    Args:
        query: Words to look for, or a date like '2026-10-25'.
        limit: Maximum number of matches to return.
    """
    return history.format_rows(history.search(query, limit))


# ======================================================================
# CONTACTS + CALLS
# A laptop has no SIM card, so Arjun cannot place a call by himself.
# What he does: open your phone dialer with the number filled in. On Windows,
# if you link your Android phone with the "Phone Link" app, the call is placed
# through your phone. You always get a yes/no question first.
# ======================================================================
CONTACTS_FILE = config.DATA_DIR / "contacts.json"
NUMBER_PATTERN = re.compile(r"^\+?[0-9][0-9 \-]{5,17}$")


def add_contact(name: str, number: str) -> str:
    """Save a contact (name and phone number) so Arjun can call them later.

    Args:
        name: The person's name.
        number: Phone number, with country code if possible, like +919876543210.
    """
    if not NUMBER_PATTERN.match(number.strip()):
        return "That doesn't look like a valid phone number."
    contacts = _load_json(CONTACTS_FILE, {})
    contacts[name.strip().lower()] = number.strip()
    _save_json(CONTACTS_FILE, contacts)
    return f"Saved {name}: {number}"


def call_contact(name_or_number: str) -> str:
    """Start a phone call to a saved contact or a phone number. The user is asked to confirm first.

    Args:
        name_or_number: A saved contact's name, or a phone number.
    """
    contacts = _load_json(CONTACTS_FILE, {})
    key = name_or_number.strip().lower()
    number = contacts.get(key, name_or_number.strip())
    if not NUMBER_PATTERN.match(number):
        return f"I don't have a contact called '{name_or_number}'. Add them first with their number."
    answer = input(f"\n   Arjun wants to call {name_or_number} ({number}). Go ahead? [y/n]: ")
    if answer.strip().lower() not in ("y", "yes"):
        return "Call cancelled by the user."
    uri = "tel:" + number.replace(" ", "").replace("-", "")
    if hasattr(os, "startfile"):
        os.startfile(uri)          # Windows: opens Phone Link / default dialer
    else:
        webbrowser.open(uri)
    return f"Opened the dialer for {name_or_number}. Press the call button to connect."


# ======================================================================
# DOCUMENTS: Word, PDF, PowerPoint
# The AI writes the text; these functions turn it into real files.
# Text format the AI uses:   "# Heading"   "## Sub-heading"   "- bullet"   plain lines = paragraphs
# For PowerPoint, every "# Heading" starts a new slide.
# ======================================================================
def _parse_content(content: str):
    blocks = []
    for line in content.splitlines():
        s = line.strip()
        if not s:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            blocks.append(("h1" if len(m.group(1)) == 1 else "h2", m.group(2)))
        elif s[:2] in ("- ", "* ", "• "):
            blocks.append(("bullet", s[2:].strip()))
        else:
            blocks.append(("p", s))
    return blocks


def _safe_filename(title: str, ext: str) -> str:
    base = re.sub(r"[^\w\- ]+", "", title, flags=re.UNICODE).strip().replace(" ", "_") or "document"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    return f"{base[:40]}_{stamp}.{ext}"


def _make_word(title, blocks, path):
    from docx import Document
    doc = Document()
    doc.add_heading(title, 0)
    for kind, text in blocks:
        if kind == "h1":
            doc.add_heading(text, 1)
        elif kind == "h2":
            doc.add_heading(text, 2)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.save(path)


def _make_ppt(title, blocks, path):
    from pptx import Presentation
    from pptx.util import Pt
    # Group the blocks into slides: each "h1" starts a new slide.
    slides = []
    for kind, text in blocks:
        if kind == "h1" or not slides:
            slides.append([text if kind == "h1" else "Overview", []])
            if kind == "h1":
                continue
        slides[-1][1].append(text)
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    first.placeholders[1].text = "Made with Arjun"
    for slide_title, lines in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = slide_title
        tf = s.placeholders[1].text_frame
        tf.word_wrap = True
        size = Pt(26) if len(lines) <= 4 else Pt(20) if len(lines) <= 7 else Pt(16)
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.size = size
    prs.save(path)


def _find_font(names):
    windir = os.environ.get("WINDIR", r"C:\Windows")
    folders = [os.path.join(windir, "Fonts"), "/usr/share/fonts/truetype/dejavu"]
    for folder in folders:
        for n in names:
            p = os.path.join(folder, n)
            if os.path.exists(p):
                return p
    return None


def _make_pdf(title, blocks, path):
    from fpdf import FPDF
    regular = _find_font(["Nirmala.ttf", "nirmala.ttf", "DejaVuSans.ttf"])
    bold = _find_font(["NirmalaB.ttf", "nirmalab.ttf", "DejaVuSans-Bold.ttf"])
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    if regular:
        pdf.add_font("ArjunFont", "", regular)
        if bold:
            pdf.add_font("ArjunFont", "B", bold)
        family, bullet = "ArjunFont", "\u2022  "
        try:
            pdf.set_text_shaping(True)   # makes Hindi/Malayalam letters join correctly
        except Exception:
            pass
    else:
        family, bullet = "Helvetica", "-  "
    style = "B" if (regular and bold) or not regular else ""

    def write(text, size, st="", indent=0):
        pdf.set_font(family, st if (st == "" or style == "B") else "", size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, size * 0.5 + 2, text, new_x="LMARGIN", new_y="NEXT")

    write(title, 22, "B")
    pdf.ln(4)
    for kind, text in blocks:
        if kind == "h1":
            pdf.ln(3); write(text, 16, "B")
        elif kind == "h2":
            pdf.ln(2); write(text, 13, "B")
        elif kind == "bullet":
            write(bullet + text, 12, "", indent=5)
        else:
            write(text, 12)
            pdf.ln(1)
    pdf.output(path)


def create_document(file_type: str, title: str, content: str) -> str:
    """Create a Word, PDF or PowerPoint file and save it in the Arjun files folder.
    Write the FULL content yourself before calling this.

    Args:
        file_type: One of 'word', 'pdf' or 'ppt'.
        title: Title of the document.
        content: The full text. Use '# Heading' for sections (each one becomes a new slide in ppt), '## Sub-heading', '- ' for bullet points, and plain lines for paragraphs.
    """
    kind = file_type.strip().lower()
    makers = {"word": ("docx", _make_word), "docx": ("docx", _make_word),
              "pdf": ("pdf", _make_pdf),
              "ppt": ("pptx", _make_ppt), "pptx": ("pptx", _make_ppt), "powerpoint": ("pptx", _make_ppt)}
    if kind not in makers:
        return "file_type must be 'word', 'pdf' or 'ppt'."
    ext, maker = makers[kind]
    config.FILES_DIR.mkdir(parents=True, exist_ok=True)
    path = str(config.FILES_DIR / _safe_filename(title, ext))
    try:
        maker(title, _parse_content(content), path)
    except ImportError as e:
        return f"A needed library is missing ({e}). Run: pip install -r requirements.txt"
    except Exception as e:
        return f"Could not create the file: {e}"
    return f"Created {ext.upper()} file: {path}"


# ======================================================================
# MUSIC (Spotify)
# play_music finds the song through the Spotify Web API and starts it in the
# Spotify desktop app by itself. This needs Spotify Premium plus a free
# developer "app" (README -> Spotify setup). Without that, Arjun falls back to
# opening the search in the Spotify app and you press play yourself.
# control_music presses the keyboard media keys (pause, next, volume...).
# ======================================================================
def _spotify_client():
    import spotipy
    from spotipy.oauth2 import SpotifyPKCE
    from spotipy.cache_handler import CacheFileHandler
    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    auth = SpotifyPKCE(
        client_id=config.SPOTIFY_CLIENT_ID,
        redirect_uri=config.SPOTIFY_REDIRECT_URI,
        scope="user-modify-playback-state user-read-playback-state",
        cache_handler=CacheFileHandler(cache_path=str(config.DATA_DIR / ".spotify_token")),
        open_browser=True,
    )
    return spotipy.Spotify(auth_manager=auth)


def _open_spotify_app(target="spotify:"):
    if hasattr(os, "startfile"):
        os.startfile(target)          # Windows: opens the Spotify desktop app
    else:
        webbrowser.open(target)


def _spotify_fallback(song, reason):
    try:
        _open_spotify_app("spotify:search:" + quote(song))
    except Exception:
        pass
    return (f"{reason} Opened the search for '{song}' in the Spotify app instead - "
            f"the user has to press play.")


def _get_device(sp):
    """Find the Spotify app to play on. Starts the app if it is closed."""
    for attempt in range(12):
        devices = sp.devices().get("devices", [])
        if devices:
            for d in devices:
                if d.get("is_active"):
                    return d
            for d in devices:
                if d.get("type") == "Computer":
                    return d
            return devices[0]
        if attempt == 0:
            _open_spotify_app()
        time.sleep(1)
    return None


def play_music(song: str) -> str:
    """Open the Spotify app and start playing a song right now. Use this whenever the user asks to play a song, artist or music. Never just give a link.

    Args:
        song: Song name, optionally with artist, e.g. 'Kesariya Arijit Singh'.
    """
    if not config.SPOTIFY_CLIENT_ID:
        return _spotify_fallback(song, "Automatic Spotify playing is not set up yet.")
    try:
        import spotipy
    except ImportError:
        return _spotify_fallback(song, "The spotipy library is missing (pip install spotipy).")
    try:
        sp = _spotify_client()
        items = sp.search(q=song, type="track", limit=1)["tracks"]["items"]
        if not items:
            return f"Spotify has no track matching '{song}'."
        track = items[0]
        label = f"{track['name']} by {', '.join(a['name'] for a in track['artists'])}"
        device = _get_device(sp)
        if device is None:
            return _spotify_fallback(song, "Could not find the Spotify app to play on.")
        sp.start_playback(device_id=device["id"], uris=[track["uri"]])
        return f"Playing {label} on Spotify."
    except spotipy.SpotifyException as e:
        if e.http_status == 403:
            return _spotify_fallback(song, "Spotify refused to play (automatic playing needs Spotify Premium).")
        return _spotify_fallback(song, f"Spotify error: {e}")
    except Exception as e:
        return _spotify_fallback(song, f"Spotify problem: {e}")


def control_music(action: str) -> str:
    """Control the music that is playing: pause, resume, next song, previous song, volume or mute.

    Args:
        action: One of 'pause', 'play', 'next', 'previous', 'volume_up', 'volume_down', 'mute'.
    """
    keys = {"pause": 0xB3, "play": 0xB3, "next": 0xB0, "previous": 0xB1,
            "volume_up": 0xAF, "volume_down": 0xAE, "mute": 0xAD}
    a = action.strip().lower().replace(" ", "_")
    if a not in keys:
        return "action must be one of: " + ", ".join(keys)
    try:
        import ctypes
        user32 = ctypes.windll.user32
    except (AttributeError, OSError):
        return "Music control only works on Windows."
    for _ in range(5 if a.startswith("volume") else 1):   # each volume press is about 2%
        user32.keybd_event(keys[a], 0, 0, 0)    # key down
        user32.keybd_event(keys[a], 0, 2, 0)    # key up
    return f"Done: {a}."


TOOLS = [add_task, list_tasks, complete_task, delete_task, search_web, get_time,
         search_history, add_contact, call_contact, create_document,
         play_music, control_music]
TOOL_LOOKUP = {fn.__name__: fn for fn in TOOLS}
